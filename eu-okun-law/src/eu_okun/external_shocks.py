"""Externally identified shock series: download, freeze, parse and aggregate.

Two published series (``IDENTIFICATION_PROTOCOL.md``, section 2), both CC BY 4.0:

- Jarociński and Karadi (2020), updated: ECB monetary policy and central bank
  information shocks from high-frequency surprises, monthly (GitHub
  ``marekjarocinski/jkshocks_update_ecb``);
- Känzig (2021), updated: oil supply surprises around OPEC announcements and the
  VAR-based oil supply news shock, monthly (GitHub ``dkaenzig/oilsupplynews``).

Files are downloaded from pinned commits, checked against SHA-256 hashes recorded at
the first download, frozen in a snapshot directory with a manifest (URL, commit,
licence, citation, hash) and aggregated to quarters by summing the three months of each
complete quarter.
"""

from __future__ import annotations

import csv
import hashlib
import json
import urllib.request
from collections.abc import Mapping
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from eu_okun.contract import EurostatDataError


@dataclass(frozen=True)
class ShockSource:
    """A pinned, published shock file."""

    key: str
    file: str
    url: str
    repository: str
    commit: str
    licence: str
    citation: str


JK_COMMIT = "f7ffc821b0ade71dd38539e4044e25368dfb4dc1"
OIL_COMMIT = "3496b6e848ba760d4be25289fc68c1e989998790"
SOURCES: dict[str, ShockSource] = {
    "monetary": ShockSource(
        "monetary", "shocks_ecb_mpd_me_m.csv",
        f"https://raw.githubusercontent.com/marekjarocinski/jkshocks_update_ecb/{JK_COMMIT}/"
        "shocks_ecb_mpd_me_m.csv",
        "https://github.com/marekjarocinski/jkshocks_update_ecb", JK_COMMIT, "CC BY 4.0",
        "Jarocinski, M. and Karadi, P. (2020). Deconstructing Monetary Policy Surprises - "
        "The Role of Information Shocks. AEJ: Macroeconomics 12(2). "
        "doi:10.1257/mac.20180090"),
    "oil": ShockSource(
        "oil", "oilSupplyNewsShocks_2025M12.xlsx",
        f"https://raw.githubusercontent.com/dkaenzig/oilsupplynews/{OIL_COMMIT}/"
        "oilSupplyNewsShocks_2025M12.xlsx",
        "https://github.com/dkaenzig/oilsupplynews", OIL_COMMIT, "CC BY 4.0",
        "Kaenzig, D. R. (2021). The Macroeconomic Effects of Oil Supply News: Evidence "
        "from OPEC Announcements. American Economic Review 111(4). "
        "doi:10.1257/aer.20190964"),
}
MONETARY_SERIES = ("MP_median", "MP_pm", "CBI_median", "CBI_pm", "pc1_hf")
OIL_SERIES = ("oil_surprise", "oil_news_shock")
QUARTERLY_FILE = "quarterly_shocks.csv"

Monthly = dict[tuple[int, int], float]


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ----------------------------------------------------------------------------------
# Parsing
# ----------------------------------------------------------------------------------


def read_monetary(path: Path) -> dict[str, Monthly]:
    """Monthly ECB shock components from the Jarociński-Karadi CSV."""
    out: dict[str, Monthly] = {name: {} for name in MONETARY_SERIES}
    with path.open(encoding="utf-8") as stream:
        reader = csv.DictReader(stream)
        missing = [n for n in ("year", "month", *MONETARY_SERIES) if n not in (reader.fieldnames
                                                                              or [])]
        if missing:
            raise EurostatDataError(f"Monetary shock file lacks columns: {', '.join(missing)}.")
        for row in reader:
            key = (int(row["year"]), int(row["month"]))
            for name in MONETARY_SERIES:
                out[name][key] = float(row[name])
    return out


def read_oil(path: Path) -> dict[str, Monthly]:
    """Monthly oil supply surprises and news shocks from the Känzig workbook."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Monthly" not in workbook.sheetnames:
        raise EurostatDataError("Oil shock workbook has no 'Monthly' sheet.")
    rows = list(workbook["Monthly"].iter_rows(values_only=True))
    workbook.close()
    header = [str(c) for c in rows[0]]
    expected = ["Date", "Oil supply surprise series", "Oil supply news shock"]
    if header[:3] != expected:
        raise EurostatDataError(f"Unexpected oil shock columns: {header[:3]}.")
    out: dict[str, Monthly] = {name: {} for name in OIL_SERIES}
    for date, surprise, shock in (r[:3] for r in rows[1:]):
        if date is None:
            continue
        year, month = str(date).split("M")
        key = (int(year), int(month))
        out["oil_surprise"][key] = float(surprise)
        out["oil_news_shock"][key] = float(shock)
    return out


def to_quarterly(monthly: Monthly) -> dict[str, float]:
    """Sum of the three months of each quarter; incomplete quarters are dropped."""
    quarters: dict[str, list[float]] = {}
    for (year, month), value in monthly.items():
        if not 1 <= month <= 12:
            raise ValueError(f"Invalid month {month}.")
        quarters.setdefault(f"{year}-Q{(month - 1) // 3 + 1}", []).append(value)
    return {q: float(sum(v)) for q, v in sorted(quarters.items()) if len(v) == 3}


def quarterly_series(directory: Path) -> dict[str, dict[str, float]]:
    """Every quarterly shock series derived from the raw files of a snapshot."""
    raw = directory / "raw"
    monthly = {**read_monetary(raw / SOURCES["monetary"].file),
               **read_oil(raw / SOURCES["oil"].file)}
    return {name: to_quarterly(values) for name, values in monthly.items()}


def _write_quarterly(path: Path, series: Mapping[str, Mapping[str, float]]) -> None:
    quarters = sorted({q for values in series.values() for q in values})
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(["quarter", *series])
        for q in quarters:
            writer.writerow([q, *(repr(values[q]) if q in values else "" for values in
                                  series.values())])


def load_quarterly(directory: Path) -> dict[str, dict[str, float]]:
    """Read ``quarterly_shocks.csv`` of a verified snapshot."""
    out: dict[str, dict[str, float]] = {}
    with (directory / QUARTERLY_FILE).open(encoding="utf-8") as stream:
        reader = csv.DictReader(stream)
        names = [n for n in (reader.fieldnames or []) if n != "quarter"]
        out = {n: {} for n in names}
        for row in reader:
            for n in names:
                if row[n] != "":
                    out[n][row["quarter"]] = float(row[n])
    return out


# ----------------------------------------------------------------------------------
# Snapshot
# ----------------------------------------------------------------------------------


def fetch_shocks(output_dir: Path, expected_sha256: Mapping[str, str] | None = None,
                 ) -> dict[str, Any]:
    """Download the pinned files, verify or record their hashes, write the snapshot."""
    raw = output_dir / "raw"
    raw.mkdir(parents=True, exist_ok=True)
    files = []
    for source in SOURCES.values():
        target = raw / source.file
        with urllib.request.urlopen(source.url, timeout=120) as response:
            target.write_bytes(response.read())
        digest = _sha256(target)
        if expected_sha256 and expected_sha256.get(source.key) not in (None, digest):
            raise EurostatDataError(f"{source.file}: SHA-256 {digest} does not match the "
                                    "recorded hash.")
        files.append({**asdict(source), "sha256": digest, "bytes": target.stat().st_size})
    series = quarterly_series(output_dir)
    _write_quarterly(output_dir / QUARTERLY_FILE, series)
    manifest = {
        "description": "Externally identified shock series for IDENTIFICATION_PROTOCOL.md",
        "files": files,
        "quarterly_csv_sha256": _sha256(output_dir / QUARTERLY_FILE),
        "aggregation": "sum of the three months of each complete quarter",
        "coverage": {name: [min(v), max(v)] for name, v in series.items()},
    }
    (output_dir / "manifest.json").write_text(json.dumps(manifest, indent=1) + "\n",
                                              encoding="utf-8", newline="\n")
    return manifest


def audit_shock_snapshot(directory: Path) -> dict[str, Any]:
    """Check the raw hashes and that the quarterly CSV re-derives from the raw files."""
    manifest: dict[str, Any] = json.loads((directory / "manifest.json").read_text(
        encoding="utf-8"))
    for entry in manifest["files"]:
        digest = _sha256(directory / "raw" / entry["file"])
        if digest != entry["sha256"]:
            raise EurostatDataError(f"{entry['file']}: SHA-256 mismatch ({digest}).")
    if _sha256(directory / QUARTERLY_FILE) != manifest["quarterly_csv_sha256"]:
        raise EurostatDataError("Quarterly shock CSV does not match its manifest hash.")
    derived = quarterly_series(directory)
    stored = load_quarterly(directory)
    for name, values in derived.items():
        if stored.get(name) != values:
            raise EurostatDataError(f"Quarterly series {name} does not re-derive from raw.")
    return manifest
