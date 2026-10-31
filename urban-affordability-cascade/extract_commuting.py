"""Audit a pinned Census-derived workbook and extract AMP commuting links.

The workbook is a THIRD-PARTY export attributed to INE Censos 2021, indicator
0012340. The INE API should be checked independently when available. Its
municipal interaction counts are symmetric: they do not give directed moves.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from pathlib import Path
from typing import Iterable, Sequence


ROOT = Path(__file__).resolve().parent
PANEL = ROOT / "data" / "amp_municipal_panel_2019_2025.csv"
OUTPUT = ROOT / "data" / "amp_census2021_commuting_pairs.csv"
PROVENANCE = ROOT / "data" / "amp_census2021_commuting_provenance.json"
SOURCE_REPO = "andrebrito0/OD_Data_Portugal_Census_2021"
SOURCE_COMMIT = "a3bd023b8eab56962cb377d1f07c6279cde67007"
SOURCE_SHA256 = "7d6dbb58be8356c868c9d0d548a81db27025228f662c5c09d71cd4ba1ee7dace"
SOURCE_URL = (
    "https://raw.githubusercontent.com/" + SOURCE_REPO + "/" + SOURCE_COMMIT + "/all_data.xlsx"
)
INDICATOR_URL = "https://www.ine.pt/xurl/indx/0012340/PT"
FIELDS = (
    "municipality_a_code", "municipality_a", "municipality_b_code",
    "municipality_b", "commuting_interactions_2021",
)
ORIGIN_RE = re.compile(r"^(\d{4}):\s*(.+)$")


def municipality_codes_from_panel(panel_path: Path = PANEL) -> dict[str, str]:
    """Use the study's fixed DICO identities; never infer codes from names."""
    from analysis import MUNICIPALITIES

    code_to_name: dict[str, str] = {}
    with panel_path.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            code = row["dico"]
            name = row["municipality"]
            if not re.fullmatch(r"\d{4}", code):
                raise ValueError(f"Invalid municipality DICO code: {code!r}")
            if code in code_to_name and code_to_name[code] != name:
                raise ValueError(f"Conflicting DICO name for {code}")
            code_to_name[code] = name
    if set(code_to_name.values()) != MUNICIPALITIES or len(code_to_name) != 17:
        raise ValueError("Panel does not contain the exact 17 AMP municipality identities")
    return code_to_name


def parse_workbook_rows(
    rows: Iterable[Sequence[object]], codes: dict[str, str]
) -> list[dict[str, str | int]]:
    """Forward-fill grouped Excel origins, then require 17x17 exact symmetry.

    INE-style exports repeat an origin once, followed by blank origin cells.
    Off-diagonal duplicate appearances are the two views of ONE undirected
    municipal interaction, and must never be added together.
    """
    iterator = iter(rows)
    header = next(iterator, None)
    if header is None or tuple(header[:3]) != ("Origem", "Destino", "n"):
        raise ValueError("Expected workbook columns Origem, Destino, n")

    observed: dict[tuple[str, str], int] = {}
    current_code: str | None = None
    inverse_codes = {name: code for code, name in codes.items()}
    for line_number, row in enumerate(iterator, start=2):
        if len(row) < 3:
            raise ValueError(f"Short workbook row {line_number}")
        origin_raw, destination_raw, count_raw = row[:3]
        if origin_raw is not None and str(origin_raw).strip():
            matched = ORIGIN_RE.fullmatch(str(origin_raw).strip())
            if matched is None:
                raise ValueError(f"Malformed origin in row {line_number}")
            current_code, source_name = matched.groups()
            if current_code in codes and source_name != codes[current_code]:
                raise ValueError(f"Origin code/name mismatch in row {line_number}")
        if current_code is None:
            raise ValueError("First source row has no origin")
        if current_code not in codes:
            continue
        destination = str(destination_raw).strip()
        if destination not in inverse_codes:
            continue
        dest_code = inverse_codes[destination]
        if isinstance(count_raw, bool) or not isinstance(count_raw, (int, float)):
            raise ValueError(f"Invalid interaction count in row {line_number}")
        if not float(count_raw).is_integer() or count_raw < 0:
            raise ValueError(f"Invalid nonnegative integer in row {line_number}")
        key = (current_code, dest_code)
        if key in observed:
            raise ValueError(f"Duplicate municipality pair in row {line_number}")
        observed[key] = int(count_raw)

    expected = {(a, b) for a in codes for b in codes}
    if set(observed) != expected:
        raise ValueError(f"Incomplete 17x17 matrix: missing {len(expected - set(observed))} cells")
    for a in codes:
        if observed[a, a] != 0:
            raise ValueError(f"Nonzero diagonal at {a}; source semantics changed")
        for b in codes:
            if observed[a, b] != observed[b, a]:
                raise ValueError(f"Asymmetric interaction counts for {a} and {b}")

    pair_rows: list[dict[str, str | int]] = []
    for i, a in enumerate(sorted(codes)):
        for b in sorted(codes)[i + 1 :]:
            pair_rows.append({
                "municipality_a_code": a,
                "municipality_a": codes[a],
                "municipality_b_code": b,
                "municipality_b": codes[b],
                "commuting_interactions_2021": observed[a, b],
            })
    return pair_rows


def extract(source: Path, *, output: Path = OUTPUT, provenance: Path = PROVENANCE) -> None:
    """Create deterministic derived data, refusing an unverified raw workbook."""
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    if digest != SOURCE_SHA256:
        raise ValueError("Upstream workbook SHA-256 does not match pinned source")
    from openpyxl import load_workbook  # Optional dependency for the extraction step.

    book = load_workbook(source, read_only=True, data_only=True)
    try:
        if len(book.worksheets) != 1:
            raise ValueError("Unexpected number of source worksheets")
        rows = parse_workbook_rows(book.worksheets[0].iter_rows(values_only=True), municipality_codes_from_panel())
    finally:
        book.close()
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    record = {
        "indicator": "INE Censos 2021 0012340 (NUTS 2024 municipality geography)",
        "indicator_url": INDICATOR_URL,
        "source_type": "third-party workbook export attributed to INE; not independently re-fetched from INE",
        "source_repository": SOURCE_REPO,
        "source_commit": SOURCE_COMMIT,
        "source_url": SOURCE_URL,
        "source_workbook_sha256": SOURCE_SHA256,
        "derived_csv_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "geography": "17 AMP municipalities, 2021",
        "row_count": len(rows),
        "semantics": "undirected symmetric pair interactions for work/study; NOT residential relocation",
        "diagonal": "all 17 self-pairs zero in source export; not a within-city commuter count",
        "license": "Statistics Portugal open-data attribution CC BY 4.0; credit INE and workbook author",
    }
    provenance.write_text(json.dumps(record, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main(argv: Sequence[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-xlsx", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--provenance", type=Path, default=PROVENANCE)
    args = parser.parse_args(argv)
    extract(args.source_xlsx, output=args.output, provenance=args.provenance)


if __name__ == "__main__":
    main()
